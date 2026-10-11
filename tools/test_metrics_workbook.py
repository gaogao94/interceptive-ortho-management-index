# -*- coding: utf-8 -*-
"""公式验收：填一组答案 → 重算 → 逐项断言公式结果。

做法：
  1. 用真实答卷（康桥样例）作为输入，验证「填答表」的逐题换算、维度小计、
     总分、百分制、等级五层公式。
  2. 期望值**手算写死**在断言里，不引用工作簿自身的中间结果 —— 否则公式错了
     自己也算不出来，测试会跟着错。

用法：python test_metrics_workbook_formulas.py <工作簿.xlsx>
退出码：0 全部通过；1 有断言失败；2 用法错误
"""
from __future__ import annotations

import os
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

NODE = r"D:\Program\DSH-desktop\resources\runtime\primary-runtime\dependencies\node\bin\node.exe"
CLI = r"D:\Program\DSH-desktop\resources\app.asar.unpacked\dsh\node_modules\@deepseek-ai\libreoffice-kit\lib\cli.js"

#: 康桥样例（case1）的 24 题判分，**从样例报告直接抄的**。
#: 不要手写这份表 —— 手写必然与样例漂移（本轮就写错过一次：B3 抄成 2 分、C2 抄成 2 分）。
FIXTURE_SCORES = {
    "A1": 1, "A2": 1, "A3": 1, "A4": 1, "A5": 1, "A6": 1,   # 6
    "B1": 1, "B2": 1, "B3": 1, "B4": 0, "B5": 1, "B6": 1,   # 5
    "C1": 2, "C2": 1, "C3": 2, "C4": 0, "C5": 1, "C6": 1,   # 7
    "D1": 1, "D2": 0, "D3": 1, "D4": 0, "D5": 1, "D6": 0,   # 3
}
#: 选项号 = 分值反查：3 分→选项 1，2 分→选项 2，1 分→选项 3，0 分→选项 4
SCORE_TO_OPTION = {3: 1, 2: 2, 1: 3, 0: 4}
ANSWERS = {q: SCORE_TO_OPTION[s] for q, s in FIXTURE_SCORES.items()}
EXPECT_SCORE = dict(FIXTURE_SCORES)
#: 手算期望（不是从工作簿读的）
EXPECT_SUB = {"A": 6, "B": 5, "C": 7, "D": 3}
EXPECT_TOTAL = 21                      # 6+5+7+3
EXPECT_PCT = round(21 / 72 * 100)      # 29
EXPECT_LEVEL = "L1 无标准"              # 29 → 0–30
#: 与报告对照：维度百分制
EXPECT_DIM_PCT = {"A": 33, "B": 28, "C": 39, "D": 17}


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print(__doc__)
        return 2
    src = argv[1]

    # 1) 填答案
    wb = load_workbook(src)
    ws = wb["填答表"]
    row_of = {}
    for r in range(3, ws.max_row + 1):
        q = ws.cell(r, 2).value
        if isinstance(q, str) and len(q) == 2 and q[0] in "ABCD" and q[1].isdigit():
            row_of[q] = r
    if len(row_of) != 24:
        print(f"  ✗ 填答表应有 24 题，实际 {len(row_of)}")
        return 1
    for q, opt in ANSWERS.items():
        ws.cell(row_of[q], 4, opt)
    work = os.path.join(tempfile.mkdtemp(prefix="iomi_xlsx_"), "filled.xlsx")
    wb.save(work)

    # 2) 重算（openpyxl 不算公式，必须过一遍重算引擎）
    recalc = work.replace("filled.xlsx", "recalc.xlsx")
    r = subprocess.run([NODE, CLI, "recalculate", "--input", work, "--output", recalc],
                       capture_output=True, text=True, encoding="utf-8")
    if r.returncode != 0 or not os.path.exists(recalc):
        print(f"  ✗ 重算失败 exit={r.returncode} {(r.stderr or '')[:200]}")
        return 1

    sf = load_workbook(recalc, data_only=False)["填答表"]   # 公式
    sv = load_workbook(recalc, data_only=True)["填答表"]    # 计算值
    fails = []

    # 3) 逐题
    bad = []
    for q, exp in EXPECT_SCORE.items():
        got = sv.cell(row_of[q], 5).value
        if got != exp:
            bad.append(f"{q}: 得 {got}，期望 {exp}")
    print(f"  逐题换算 24 题：{'全部正确' if not bad else '错误 ' + str(bad)}")
    if bad:
        fails.append("逐题换算")

    # 4) 维度小计 / 总分 / 百分制 / 等级
    labels = {}
    for r2 in range(3, sf.max_row + 1):
        lab = sf.cell(r2, 2).value
        if isinstance(lab, str):
            labels[lab] = r2

    for dim, exp in EXPECT_SUB.items():
        r2 = labels.get(f"{dim} 小计")
        if r2 is None:
            fails.append(f"缺 {dim} 小计 行")
            continue
        got = sv.cell(r2, 5).value
        ok = got == exp
        pct = round(exp / 18 * 100)
        print(f"  {dim} 小计：公式值 {got}，手算 {exp}（折百分制 {pct}，报告写 {EXPECT_DIM_PCT[dim]}）"
              f"  {'OK' if ok and pct == EXPECT_DIM_PCT[dim] else '✗'}")
        if not ok:
            fails.append(f"{dim} 小计")
        if pct != EXPECT_DIM_PCT[dim]:
            fails.append(f"{dim} 折百分制与报告不符")

    for lab, exp in ((f"四维合计 / 72", EXPECT_TOTAL),
                     ("合计 ÷ 72 × 100", EXPECT_PCT),
                     ("对照区间", EXPECT_LEVEL)):
        r2 = labels.get(lab)
        if r2 is None:
            fails.append(f"缺 {lab} 行")
            continue
        got = sv.cell(r2, 5).value
        ok = got == exp
        print(f"  {lab}：公式值 {got!r}，手算 {exp!r}  {'OK' if ok else '✗'}")
        if not ok:
            fails.append(lab)

    # 5) 公式必须还在（不能被值替换）
    formulas = sum(1 for row in sf.iter_rows() for c in row
                   if isinstance(c.value, str) and c.value.startswith("="))
    print(f"  工作簿里保留的公式数：{formulas}")
    if formulas < 26:
        fails.append("公式被值替换")

    if fails:
        print(f"\n  ✗ 失败项：{fails}")
        return 1
    print("\n  ✓ 公式全部正确（逐题 24 + 小计 4 + 总分 2 + 等级 1）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
