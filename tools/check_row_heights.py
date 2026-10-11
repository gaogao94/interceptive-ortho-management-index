# -*- coding: utf-8 -*-
"""核对行高：每一行都必须显式设高，且不低于内容所需。

判定标准：
  - 有内容的行必须设了行高（程序生成的工作簿里自动行高不生效）
  - 行高 >= 按内容估算的所需高度（与生成脚本同一套估算法）

用法：python check_row_heights.py <工作簿.xlsx>
退出码：0 通过；1 有问题；2 用法错误
"""
from __future__ import annotations

import math
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

LINE_PT = 14.0
CLIP_PT = 8.0
MIN_ROW_PT = 22.0


def needed(values: list, widths: list[float], merged_span: int = 0) -> float:
    """估算所需行高。merged_span > 0 表示这是合并行，按合并后的总宽算。"""
    if merged_span:
        widths = [sum(widths[:merged_span])]
        values = values[:1]
    lines = 1
    for text, width in zip(values, widths):
        if text is None:
            continue
        w = max(float(width), 4.0)
        n = 0
        for part in str(text).split("\n"):
            n += max(1, math.ceil(len(part) / (w * 1.0)))
        lines = max(lines, n)
    return max(MIN_ROW_PT, lines * LINE_PT + CLIP_PT)


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print(__doc__)
        return 2
    wb = load_workbook(argv[1])

    fails = []
    for ws in wb.worksheets:
        widths = []
        for c in range(1, ws.max_column + 1):
            d = ws.column_dimensions.get(get_column_letter(c))
            widths.append(d.width if d and d.width else 8.43)

        print(f"\n【{ws.title}】{ws.max_row} 行 × {ws.max_column} 列　列宽 {[round(w,1) for w in widths]}")
        # 合并区域的左上角单元格按合并总宽算
        merge_span = {}
        for m in ws.merged_cells.ranges:
            merge_span[(m.min_row, m.min_col)] = m.max_col - m.min_col + 1

        no_height, too_short = [], []
        for r in range(1, ws.max_row + 1):
            cells = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            if not any(v is not None for v in cells):
                continue
            h = ws.row_dimensions[r].height if r in ws.row_dimensions else None
            if h is None:
                no_height.append(r)
                continue
            span = merge_span.get((r, 1), 0)
            need = needed(cells, widths, merged_span=span)
            if h + 0.5 < need:
                too_short.append((r, round(h, 1), round(need, 1)))

        print(f"  有内容但未设行高: {no_height if no_height else '无'}")
        print(f"  行高不足的行:     {too_short if too_short else '无'}")
        if no_height:
            fails.append(f"{ws.title} 未设行高 {len(no_height)} 行")
        if too_short:
            fails.append(f"{ws.title} 行高不足 {len(too_short)} 行")

        hs = [ws.row_dimensions[r].height for r in range(1, ws.max_row + 1)
              if r in ws.row_dimensions and ws.row_dimensions[r].height]
        if hs:
            print(f"  行高范围: {min(hs):.1f} – {max(hs):.1f} 磅")

    print()
    if fails:
        print("✗ " + "; ".join(fails))
        return 1
    print("✓ 所有有内容的行都设了足够的行高")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
